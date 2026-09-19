import sqlite3
import datetime
import os


def generate_report():
    """Generate a plain text report from database alerts."""
    conn = sqlite3.connect("nids.db")
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM alerts")
    data = cursor.fetchall()
    conn.close()

    filename = "nids_report.txt"
    with open(filename, "w", encoding="utf-8") as f:
        f.write("NIDS SECURITY INCIDENT REPORT\n")
        f.write("=" * 50 + "\n")
        f.write(f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        f.write(f"Total Alerts: {len(data)}\n")
        f.write("=" * 50 + "\n\n")
        for row in data:
            f.write(f"ID       : {row[0]}\n")
            f.write(f"Time     : {row[1]}\n")
            f.write(f"Alert    : {row[2]}\n")
            f.write(f"Severity : {row[3]}\n")
            f.write("-" * 40 + "\n")
    print("Text report generated:", filename)
    return filename


def generate_pdf_report(filename="nids_report.pdf"):
    """Generate a styled PDF report from database alerts using ReportLab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer,
        Table, TableStyle, HRFlowable
    )

    conn = sqlite3.connect("nids.db")
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM alerts ORDER BY id DESC")
    data = cursor.fetchall()
    conn.close()

    doc = SimpleDocTemplate(
        filename,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm
    )

    styles = getSampleStyleSheet()

    # Custom Styles
    title_style = ParagraphStyle(
        "CyberTitle",
        parent=styles["Title"],
        fontSize=22,
        textColor=colors.HexColor("#00F0FF"),
        spaceAfter=6,
        fontName="Helvetica-Bold"
    )
    subtitle_style = ParagraphStyle(
        "CyberSub",
        parent=styles["Normal"],
        fontSize=10,
        textColor=colors.HexColor("#94A3B8"),
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Heading2"],
        fontSize=13,
        textColor=colors.HexColor("#FFFFFF"),
        backColor=colors.HexColor("#0C101B"),
        spaceAfter=8,
        spaceBefore=16,
        fontName="Helvetica-Bold",
        leftIndent=5
    )
    normal_style = ParagraphStyle(
        "CyberNormal",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#E2E8F0"),
    )

    elements = []

    # Header
    elements.append(Paragraph("🛡 CYBER_SHIELD // NIDS", title_style))
    elements.append(Paragraph("Network Intrusion Detection System — Security Incident Report", subtitle_style))
    elements.append(Paragraph(f"Generated: {datetime.datetime.now().strftime('%d %B %Y at %H:%M:%S')}", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#00F0FF")))
    elements.append(Spacer(1, 0.4 * cm))

    # Summary Stats
    elements.append(Paragraph("INCIDENT SUMMARY", section_style))
    total = len(data)
    critical = sum(1 for r in data if "HIGH" in str(r[3]).upper() or "CRITICAL" in str(r[2]).upper())
    medium = sum(1 for r in data if "MEDIUM" in str(r[3]).upper())
    low = sum(1 for r in data if "LOW" in str(r[3]).upper())

    summary_data = [
        ["Metric", "Count"],
        ["Total Alerts", str(total)],
        ["High / Critical", str(critical)],
        ["Medium", str(medium)],
        ["Low", str(low)],
    ]
    summary_table = Table(summary_data, colWidths=[8 * cm, 4 * cm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F1626")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#00F0FF")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#0C101B"), colors.HexColor("#080B12")]),
        ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#E2E8F0")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#1E293B")),
        ("ALIGN", (1, 0), (1, -1), "CENTER"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.5 * cm))

    # Alert Records
    elements.append(Paragraph("ALERT LOG", section_style))

    if data:
        table_header = [["#", "Timestamp", "Alert", "Severity"]]
        table_rows = []
        for row in data:
            table_rows.append([
                str(row[0]),
                str(row[1]),
                Paragraph(str(row[2]), normal_style),
                str(row[3])
            ])
        table_data = table_header + table_rows
        col_widths = [1 * cm, 4 * cm, 9 * cm, 3 * cm]
        alert_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        alert_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F1626")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#00F0FF")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#0C101B"), colors.HexColor("#080B12")]),
            ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#E2E8F0")),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#1E293B")),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("ALIGN", (3, 0), (3, -1), "CENTER"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(alert_table)
    else:
        elements.append(Paragraph("No alerts recorded yet.", normal_style))

    elements.append(Spacer(1, 1 * cm))
    elements.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#1E293B")))
    elements.append(Paragraph("CYBER_SHIELD NIDS — Confidential Security Report", subtitle_style))

    doc.build(elements)
    print("PDF report generated:", filename)
    return filename


def generate_docx_report(filename="nids_report.docx"):
    """Generate a styled Word document report from database alerts using python-docx."""
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    conn = sqlite3.connect("nids.db")
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM alerts ORDER BY id DESC")
    data = cursor.fetchall()
    conn.close()

    doc = Document()

    # Page margins
    section = doc.sections[0]
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # Helper to set paragraph color
    def set_color(run, hex_color):
        r, g, b = int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)

    # Title
    title_para = doc.add_paragraph()
    title_run = title_para.add_run("CYBER_SHIELD // NIDS — Security Incident Report")
    title_run.bold = True
    title_run.font.size = Pt(18)
    set_color(title_run, "00F0FF")
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Subtitle
    sub_para = doc.add_paragraph()
    sub_run = sub_para.add_run(
        f"Network Intrusion Detection System\nGenerated: {datetime.datetime.now().strftime('%d %B %Y at %H:%M:%S')}"
    )
    sub_run.font.size = Pt(9)
    set_color(sub_run, "94A3B8")
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # Spacer

    # Summary Section
    heading = doc.add_heading("INCIDENT SUMMARY", level=2)
    heading.runs[0].font.size = Pt(13)

    total = len(data)
    critical = sum(1 for r in data if "HIGH" in str(r[3]).upper() or "CRITICAL" in str(r[2]).upper())
    medium = sum(1 for r in data if "MEDIUM" in str(r[3]).upper())
    low = sum(1 for r in data if "LOW" in str(r[3]).upper())

    summary_table = doc.add_table(rows=5, cols=2)
    summary_table.style = "Table Grid"
    rows_data = [
        ("Metric", "Count"),
        ("Total Alerts", str(total)),
        ("High / Critical", str(critical)),
        ("Medium", str(medium)),
        ("Low", str(low)),
    ]
    for i, (label, val) in enumerate(rows_data):
        row = summary_table.rows[i]
        row.cells[0].text = label
        row.cells[1].text = val
        if i == 0:
            for cell in row.cells:
                cell.paragraphs[0].runs[0].bold = True
                cell.paragraphs[0].runs[0].font.size = Pt(10)

    doc.add_paragraph()

    # Alert Log Section
    heading2 = doc.add_heading("ALERT LOG", level=2)
    heading2.runs[0].font.size = Pt(13)

    if data:
        cols = 4
        alert_table = doc.add_table(rows=1 + len(data), cols=cols)
        alert_table.style = "Table Grid"

        # Header row
        header_row = alert_table.rows[0]
        for j, col_name in enumerate(["#", "Timestamp", "Alert", "Severity"]):
            header_row.cells[j].text = col_name
            header_row.cells[j].paragraphs[0].runs[0].bold = True
            header_row.cells[j].paragraphs[0].runs[0].font.size = Pt(10)

        # Data rows
        for i, row_data in enumerate(data):
            row = alert_table.rows[i + 1]
            row.cells[0].text = str(row_data[0])
            row.cells[1].text = str(row_data[1])
            row.cells[2].text = str(row_data[2])
            row.cells[3].text = str(row_data[3])
            for cell in row.cells:
                cell.paragraphs[0].runs[0].font.size = Pt(8)
    else:
        doc.add_paragraph("No alerts recorded yet.")

    doc.add_paragraph()
    footer_para = doc.add_paragraph("CYBER_SHIELD NIDS — Confidential Security Report")
    footer_run = footer_para.runs[0]
    footer_run.font.size = Pt(8)
    set_color(footer_run, "475569")
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(filename)
    print("DOCX report generated:", filename)
    return filename